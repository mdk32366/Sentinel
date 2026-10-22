"""Poland - Ministry of Finance wholesale bond auctions (Sentinel D-0116).

Source: the "transaction database" page of the Polish Ministry of Finance,

    https://www.gov.pl/web/finance/transaction-database

links its spreadsheets as ``/attachment/<uuid>``. The one read here is the
attachment whose Content-Disposition names ``Obligacje_Hurtowe.xlsm`` (the
uuid changes; the file name does not). Size cap 10 MB; the content type must
be xlsm, xlsx or octet-stream.

The workbook is a macro-enabled file. It is opened with openpyxl
``read_only=True, keep_vba=False, data_only=True``: macros are never loaded,
let alone run, and formulas are read as their cached values.

Sheet ``Operacje``. A row counts only when ``TypOperacji == "AS"`` (the
regular auction sale); floating-rate bonds (``Kupon == "FRN"`` or series
``NZ*``) and rows with a blank ``RentŚr`` are skipped. ``RentŚr`` is the
average auction yield as a fraction (0.06447 -> 6.447 %).

Per auction date each bond's remaining life is (DataWykupu - DataTransakcji)
/ 365.25 years. Buckets: 2Y [1.5, 3.5), 5Y [3.5, 7.0), 10Y [7.0, 12.0]; in each
the bond closest to the target tenor is taken, a tie going to the longer
bond. Stored as FB_PL_{02,05,10}Y on the auction date: "as of last auction",
not a daily curve. Display only, never scored.
"""
from __future__ import annotations

import io
import logging
import re
import time
import urllib.request
import zipfile
from datetime import datetime, timedelta
from decimal import Decimal
from typing import Callable, Optional
from urllib.parse import unquote, urljoin

from sqlalchemy.orm import Session

from pipelines.foreign_bond_common import (
    USER_AGENT,
    EmptyResponse,
    SchemaDrift,
    TIMEOUT_S,
    code_for,
    fetch_bytes,
    make_metrics,
    reject_html,
    to_value,
    write_run,
)

logger = logging.getLogger(__name__)

PIPELINE_NAME = "MinFin_PL"
HOST = "www.gov.pl"
PAGE_URL = f"https://{HOST}/web/finance/transaction-database"
FILE_NAME = "Obligacje_Hurtowe.xlsm"
SHEET = "Operacje"
MAX_BYTES = 10 * 1024 * 1024
MAX_UNZIPPED_BYTES = 200 * 1024 * 1024
CONTENT_TYPES = (
    "application/vnd.ms-excel.sheet.macroenabled.12",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/octet-stream",
)
COLUMNS = ("DataTransakcji", "DataWykupu", "Seria", "Kupon", "TypOperacji", "RentŚr")
TENORS = (2, 5, 10)
BUCKETS = {2: (1.5, 3.5, False), 5: (3.5, 7.0, False), 10: (7.0, 12.0, True)}
DAILY_DAYS = 120
BACKFILL_SINCE = datetime(2010, 1, 1)

METRICS = make_metrics(
    "PL", TENORS,
    "Poland {t}Y government bond auction yield (MinFin)",
    "Ministry of Finance (Poland)",
    "Polish Ministry of Finance Obligacje_Hurtowe.xlsm, sheet Operacje: average "
    "yield (RentSr) of the fixed-rate bond closest to {t}Y sold at a regular "
    "auction (AS), on the auction date. As of last auction, not daily. "
    "Display only, not scored (Sentinel D-0116).",
)

_ATTACH = re.compile(rb'href="(/attachment/[0-9a-fA-F-]{36})"')


def attachment_links(page: bytes) -> list:
    reject_html_ok = page.lstrip()[:1] == b"<"
    if not reject_html_ok:
        raise SchemaDrift(f"{PIPELINE_NAME}: the transaction-database page is not HTML")
    seen = []
    for m in _ATTACH.finditer(page):
        url = urljoin(PAGE_URL, m.group(1).decode("ascii"))
        if url not in seen:
            seen.append(url)
    if not seen:
        raise SchemaDrift(f"{PIPELINE_NAME}: no /attachment/ links on the page")
    return seen


def filename_of(disposition: Optional[str]) -> Optional[str]:
    if not disposition:
        return None
    m = re.search(r"filename\*\s*=\s*[^']*'[^']*'([^;]+)", disposition, re.I)
    if m:
        return unquote(m.group(1).strip().strip('"'))
    m = re.search(r'filename\s*=\s*"?([^";]+)"?', disposition, re.I)
    return m.group(1).strip() if m else None


def _head(url: str) -> dict:
    req = urllib.request.Request(url, method="HEAD", headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(req, timeout=TIMEOUT_S) as resp:
        return {k.lower(): v for k, v in resp.headers.items()}


def resolve_attachment(page: bytes, head: Callable[[str], dict] = _head) -> str:
    """The attachment URL whose Content-Disposition is Obligacje_Hurtowe.xlsm."""
    for url in attachment_links(page):
        hdrs = {k.lower(): v for k, v in (head(url) or {}).items()}
        if filename_of(hdrs.get("content-disposition")) != FILE_NAME:
            continue
        ctype = (hdrs.get("content-type") or "").split(";")[0].strip().lower()
        if ctype not in CONTENT_TYPES:
            raise SchemaDrift(f"{PIPELINE_NAME}: {FILE_NAME} has content type {ctype!r}")
        length = hdrs.get("content-length")
        if length and length.isdigit() and int(length) > MAX_BYTES:
            raise SchemaDrift(f"{PIPELINE_NAME}: {FILE_NAME} is {length} B, over the {MAX_BYTES} B cap")
        return url
    raise SchemaDrift(f"{PIPELINE_NAME}: no attachment named {FILE_NAME}")


def _check_container(body: bytes) -> None:
    if len(body) > MAX_BYTES:
        raise SchemaDrift(f"{PIPELINE_NAME}: workbook is {len(body)} B, over the {MAX_BYTES} B cap")
    reject_html(body, PIPELINE_NAME)
    if body[:4] != b"PK\x03\x04":
        raise SchemaDrift(f"{PIPELINE_NAME}: workbook is not an OOXML (zip) file")
    try:
        infos = zipfile.ZipFile(io.BytesIO(body)).infolist()
    except zipfile.BadZipFile as exc:
        raise SchemaDrift(f"{PIPELINE_NAME}: workbook zip is unreadable") from exc
    if sum(i.file_size for i in infos) > MAX_UNZIPPED_BYTES:
        raise SchemaDrift(f"{PIPELINE_NAME}: workbook unzips past {MAX_UNZIPPED_BYTES} B")


def _open_workbook(body: bytes):
    import openpyxl
    return openpyxl.load_workbook(io.BytesIO(body), read_only=True, keep_vba=False, data_only=True)


def _as_datetime(raw, where):
    if isinstance(raw, datetime):
        return raw.replace(hour=0, minute=0, second=0, microsecond=0)
    if isinstance(raw, str):
        for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return datetime.strptime(raw.strip(), fmt)
            except ValueError:
                pass
    raise SchemaDrift(f"{where}: {raw!r} is not a date")


def bucket_for(years: float):
    for tenor, (lo, hi, closed) in BUCKETS.items():
        if lo <= years < hi or (closed and years == hi):
            return tenor
    return None


def parse_auctions(body: bytes, since: Optional[datetime] = None) -> list:
    """Obligacje_Hurtowe.xlsm -> [(code, auction datetime, Decimal %)]."""
    _check_container(body)
    try:
        wb = _open_workbook(body)
    except Exception as exc:  # openpyxl raises a zoo of types on a bad file
        raise SchemaDrift(f"{PIPELINE_NAME}: workbook does not open: {type(exc).__name__}") from exc
    try:
        if SHEET not in wb.sheetnames:
            raise SchemaDrift(f"{PIPELINE_NAME}: no sheet {SHEET!r} (has {wb.sheetnames})")
        rows = wb[SHEET].iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            raise EmptyResponse(f"{PIPELINE_NAME}: sheet {SHEET} is empty")
        names = [str(h).strip() if h is not None else "" for h in header]
        missing = [c for c in COLUMNS if c not in names]
        if missing:
            raise SchemaDrift(f"{PIPELINE_NAME}: sheet {SHEET} lacks {missing}")
        ix = {c: names.index(c) for c in COLUMNS}
        # (date, tenor) -> (distance, years, value)
        best: dict = {}
        for n, r in enumerate(rows, start=2):
            if r is None or all(c is None or c == "" for c in r):
                continue
            if str(r[ix["TypOperacji"]] or "").strip() != "AS":
                continue
            seria = str(r[ix["Seria"]] or "").strip().upper()
            kupon = r[ix["Kupon"]]
            if seria.startswith("NZ") or (isinstance(kupon, str) and kupon.strip().upper() == "FRN"):
                continue
            rent = r[ix["RentŚr"]]
            if rent is None or (isinstance(rent, str) and rent.strip() == ""):
                continue
            where = f"{PIPELINE_NAME} row {n} {seria}"
            traded = _as_datetime(r[ix["DataTransakcji"]], where)
            if since is not None and traded < since:
                continue
            matures = _as_datetime(r[ix["DataWykupu"]], where)
            years = (matures - traded).days / 365.25
            tenor = bucket_for(years)
            if tenor is None:
                continue
            if isinstance(rent, str):
                try:
                    frac = Decimal(rent.strip().replace(",", "."))
                except Exception as exc:
                    raise SchemaDrift(f"{where}: RentŚr {rent!r} is not a number") from exc
            else:
                frac = Decimal(repr(float(rent)))
            value = to_value(str((frac * 100).normalize()), f"{where} RentŚr")
            if value is None:
                continue
            dist = abs(years - tenor)
            key = (traded, tenor)
            cur = best.get(key)
            # Closest to the target; a tie goes to the longer bond.
            if cur is None or dist < cur[0] or (dist == cur[0] and years > cur[1]):
                best[key] = (dist, years, value)
    finally:
        wb.close()
    out = [(code_for("PL", t), d, v[2]) for (d, t), v in best.items()]
    out.sort(key=lambda x: (x[1], x[0]))
    return out


def run_minfin_auctions(db: Session, backfill: bool = False,
                        head: Callable[[str], dict] = _head, **fetch_kw) -> dict:
    since = BACKFILL_SINCE if backfill else datetime.utcnow() - timedelta(days=DAILY_DAYS)

    def produce():
        page = fetch_bytes(PAGE_URL, headers={"Accept": "text/html"}, **fetch_kw)
        url = resolve_attachment(page, head)
        body = fetch_bytes(url, **fetch_kw)
        return parse_auctions(body, since=since)

    return write_run(db, PIPELINE_NAME, METRICS, produce, PAGE_URL)
