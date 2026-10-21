"""UK - Bank of England fitted nominal gilt spot curve (Sentinel D-0111, vault:D-0057).

Source: Bank of England yield curves, official and free.

    Daily:    https://www.bankofengland.co.uk/-/media/boe/files/statistics/
              yield-curves/latest-yield-curve-data.zip
              member "GLC Nominal daily data current month.xlsx"
    Backfill: https://www.bankofengland.co.uk/-/media/boe/files/statistics/
              yield-curves/glcnominalddata.zip
              members "GLC Nominal daily data_<from> to <to|present>.xlsx"
              (resolved from the BoE yield-curves page and fetched 2026-10-21;
              see findings.md)

Sheet ``4. spot curve``: a row labelled ``years:`` in column A carries the
maturities in 0.5-year steps; data rows carry a date in column A. Columns are
looked up by the ``years:`` value, never by index. Type FIT (fitted spot).

Lag: one business day (on 2026-10-21 the latest row was 2026-10-20). The
current-month file holds only this month, so the daily job must run every
business day. On the first business day of a month it may hold no row yet;
that run is ``failed`` (0 rows is never a success) and the badge shows it.
"""
from __future__ import annotations

import io
import logging
import re
import zipfile
from datetime import datetime, timedelta

from sqlalchemy.orm import Session

from pipelines.foreign_bond_common import (
    CATEGORY,
    TENORS,
    SchemaDrift,
    code_for,
    fetch_bytes,
    to_value,
    write_run,
)

logger = logging.getLogger(__name__)

PIPELINE_NAME = "BoE_Yields"
HOST = "www.bankofengland.co.uk"
_BASE = f"https://{HOST}/-/media/boe/files/statistics/yield-curves/"
LATEST_URL = _BASE + "latest-yield-curve-data.zip"
ARCHIVE_URL = _BASE + "glcnominalddata.zip"
LATEST_MEMBER = "GLC Nominal daily data current month.xlsx"
ARCHIVE_MEMBER_RE = re.compile(
    r"^GLC Nominal daily data_(\d{4}) to (\d{4}|present)\.xlsx$"
)
SHEET = "4. spot curve"
BACKFILL_DAYS = 1825

METRICS = [
    {
        "code": code_for("GB", t),
        "name": f"UK {t}Y gilt yield (BoE fitted spot)",
        "category": CATEGORY,
        "unit": "%",
        "source": "BoE",
        "description": (
            f"FIT. Bank of England GLC nominal spot curve, sheet '{SHEET}', "
            f"years: {float(t)}, daily, one business day lag. "
            f"latest-yield-curve-data.zip. Display only, not scored "
            f"(Sentinel D-0111)."
        ),
    }
    for t in TENORS
]


def _member(zip_bytes: bytes, name: str) -> bytes:
    try:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
            if name not in zf.namelist():
                raise SchemaDrift(f"BoE zip lacks {name!r} (has {zf.namelist()!r})")
            return zf.read(name)
    except zipfile.BadZipFile as exc:
        raise SchemaDrift(f"BoE body is not a zip: {exc}") from exc


def archive_members(zip_bytes: bytes, since_year: int) -> list:
    """Archive xlsx names whose period reaches ``since_year`` or later."""
    try:
        with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
            names = zf.namelist()
    except zipfile.BadZipFile as exc:
        raise SchemaDrift(f"BoE archive is not a zip: {exc}") from exc
    picked = []
    for name in names:
        m = ARCHIVE_MEMBER_RE.match(name)
        if not m:
            continue
        end = 9999 if m.group(2) == "present" else int(m.group(2))
        if end >= since_year:
            picked.append(name)
    if not picked:
        raise SchemaDrift(f"BoE archive has no nominal daily member since {since_year}")
    return sorted(picked)


def parse_boe_xlsx(xlsx_bytes: bytes) -> list:
    """GLC nominal xlsx -> [(code, datetime, Decimal)], oldest first."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise SchemaDrift(f"BoE xlsx unreadable: {exc}") from exc
    if SHEET not in wb.sheetnames:
        raise SchemaDrift(f"BoE xlsx lacks sheet {SHEET!r} (has {wb.sheetnames!r})")

    cols = None
    out = []
    for n, row in enumerate(wb[SHEET].iter_rows(values_only=True), start=1):
        if not row:
            continue
        first = row[0]
        if cols is None:
            if isinstance(first, str) and first.strip() == "years:":
                found = {}
                for idx, cell in enumerate(row[1:], start=1):
                    if isinstance(cell, (int, float)) and float(cell) in [float(t) for t in TENORS]:
                        found.setdefault(int(float(cell)), idx)
                missing = [t for t in TENORS if t not in found]
                if missing:
                    raise SchemaDrift(f"BoE 'years:' row lacks maturities {missing}")
                cols = found
            continue

        cells = {t: (row[i] if i < len(row) else None) for t, i in cols.items()}
        has_values = any(
            isinstance(v, (int, float)) and not isinstance(v, bool)
            for v in cells.values()
        )
        if not isinstance(first, datetime):
            # Banner rows ('#VALUE!' in the current-month file, a row of
            # 'Refresh' in the archive) and padding carry no number.
            if not has_values:
                continue
            raise SchemaDrift(f"BoE row {n}: values under an unparseable date {first!r}")
        for t in TENORS:
            code = code_for("GB", t)
            value = to_value(cells[t], f"{code} {first.date()}")
            if value is not None:
                out.append((code, datetime(first.year, first.month, first.day), value))

    if cols is None:
        raise SchemaDrift("BoE sheet has no 'years:' row")
    out.sort(key=lambda r: (r[1], r[0]))
    return out


def run_boe_yields(db: Session, backfill: bool = False, **fetch_kw) -> dict:
    def produce():
        rows = parse_boe_xlsx(_member(fetch_bytes(LATEST_URL, **fetch_kw), LATEST_MEMBER))
        if backfill:
            since = datetime.utcnow() - timedelta(days=BACKFILL_DAYS)
            archive = fetch_bytes(ARCHIVE_URL, **fetch_kw)
            for name in archive_members(archive, since.year):
                rows.extend(r for r in parse_boe_xlsx(_member(archive, name)) if r[1] >= since)
        return rows

    return write_run(db, PIPELINE_NAME, METRICS, produce, LATEST_URL)


if __name__ == "__main__":
    import sys

    logging.basicConfig(level=logging.INFO)
    from database.connection import get_session

    session = get_session()
    try:
        print(run_boe_yields(session, backfill="--backfill" in sys.argv))
    finally:
        session.close()
