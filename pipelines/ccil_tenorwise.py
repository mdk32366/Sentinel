"""India - CCIL Tenorwise Indicative Yields, daily, manual upload (Sentinel D-0118).

Source: The Clearing Corporation of India Ltd. (CCIL), Tenorwise Indicative
Yields, https://www.ccilindia.com/tenorwise-indicative-yields. The download is
sign-in gated and built client-side (F-entry, D-0118), so there is no fetcher
and no scheduler job: Matt uploads the .xlsx he downloaded, raw, to

    POST /api/ingest/ccil-tenorwise      (Basic Auth; body = the file bytes)

The file is one sheet, headers exactly ``Date | Tenor Bucket | Security |
YTM (%)``, Date a ``DD-MM-YYYY`` string, every row the same date. Four
buckets are kept, each must hold one central-government G-sec
(``^\\d+\\.\\d+% GS \\d{4}$``):

    1Y-2Y   -> FB_IN_02Y      4Y-5Y  -> FB_IN_05Y
    9Y-10Y  -> FB_IN_10Y      28Y-30Y -> FB_IN_30Y

13Y-15Y, the T-bills (91D/182D/364D) and the SDL rows (5Y/10Y/15Y) are
known and dropped. Any other bucket, a missing or duplicated kept bucket, or
any other deviation rejects the whole file (HTTP 422, nothing written).

D-0118 amendment (@0.3 §3 trimmed at build: provenance = existing row, raw
body): provenance is the existing ingest_provenance row (source, uploaded_at,
uploader, sha256, counts, min/max date); the per-row code, date and YTM live
in the yield rows; the bond names are validated, not stored; ATTRIBUTION is a
code constant the card renders. Display only, never scored.
"""
from __future__ import annotations

import hashlib
import io
import logging
import re
import zipfile
from datetime import datetime
from decimal import Decimal

from sqlalchemy.orm import Session

from pipelines.foreign_bond_common import code_for, make_metrics

logger = logging.getLogger(__name__)

PIPELINE_NAME = "CCIL_Tenorwise"
INGEST_SOURCE = "ccil-tenorwise"
ATTRIBUTION = ("Source: The Clearing Corporation of India Ltd. (CCIL), Tenorwise Indicative Yields, "
               "https://www.ccilindia.com/tenorwise-indicative-yields")
HEADERS = ("Date", "Tenor Bucket", "Security", "YTM (%)")
KEPT = {"1Y-2Y": 2, "4Y-5Y": 5, "9Y-10Y": 10, "28Y-30Y": 30}
DROPPED = {"13Y-15Y", "91D", "182D", "364D", "5Y", "10Y", "15Y"}
G_SEC = re.compile(r"^\d+\.\d+% GS \d{4}$")
DATE_FORMAT = "%d-%m-%Y"
DATE_SHAPE = re.compile(r"^\d{2}-\d{2}-\d{4}$")
YTM_RANGE = (Decimal("0"), Decimal("25"))
TENORS = (2, 5, 10, 30)
MAX_BODY_BYTES = 2 * 1024 * 1024
MAX_UNZIPPED_BYTES = 16 * 1024 * 1024
QUANT = Decimal("0.00000001")  # TimeSeries.value is Numeric(20, 8)

METRICS = make_metrics(
    "IN", TENORS,
    "India {t}Y G-sec yield (CCIL indicative)",
    "CCIL",
    "BMK. CCIL Tenorwise Indicative Yields, YTM of the benchmark central-government G-sec in the "
    "{t}Y bucket, daily, uploaded manually. Display only, not scored (Sentinel D-0118).",
)


class IngestRejected(ValueError):
    """The upload breaks the schema; HTTP 422 and nothing is written."""


def _check_container(body: bytes) -> None:
    if not body.startswith(b"PK"):
        raise IngestRejected("not an .xlsx (no ZIP signature); HTML or CSV is not accepted")
    try:
        with zipfile.ZipFile(io.BytesIO(body)) as zf:
            infos = zf.infolist()
            names = {i.filename.lower() for i in infos}
            if sum(i.file_size for i in infos) > MAX_UNZIPPED_BYTES:
                raise IngestRejected("unpacked workbook is over the size cap")
            if any(n.endswith("vbaproject.bin") for n in names):
                raise IngestRejected("workbook carries macros (.xlsm / vbaProject.bin)")
            if "[content_types].xml" in names:
                ct = zf.read("[Content_Types].xml").decode("utf-8", "replace").lower()
                if "macroenabled" in ct:
                    raise IngestRejected("workbook is macro-enabled (.xlsm)")
    except zipfile.BadZipFile as exc:
        raise IngestRejected("not a readable .xlsx") from exc


def parse_ccil(body: bytes) -> list:
    """The file -> [(code, datetime, Decimal YTM, security)] for the four kept
    buckets. Raises IngestRejected on any deviation."""
    import openpyxl

    _check_container(body)
    try:
        wb = openpyxl.load_workbook(io.BytesIO(body), read_only=True, data_only=True, keep_vba=False)
    except Exception as exc:  # openpyxl raises a zoo of types on a bad file
        raise IngestRejected(f"not a readable .xlsx: {type(exc).__name__}") from exc
    try:
        if len(wb.sheetnames) != 1:
            raise IngestRejected(f"expected 1 sheet, got {len(wb.sheetnames)}")
        rows = [r for r in wb[wb.sheetnames[0]].iter_rows(values_only=True)
                if any(c not in (None, "") for c in r)]
    finally:
        wb.close()
    if not rows:
        raise IngestRejected("empty sheet")
    header = tuple(rows[0])
    while header and header[-1] is None:
        header = header[:-1]
    if header != HEADERS:
        raise IngestRejected(f"headers {header!r} != {HEADERS!r}")
    out, seen, dates = [], set(), set()
    for n, raw in enumerate(rows[1:], start=2):
        r = tuple(raw[:4]) + (None,) * (4 - len(raw[:4]))
        if any(c not in (None, "") for c in raw[4:]):
            raise IngestRejected(f"row {n}: data beyond the four columns")
        date_s, bucket, security, ytm = r
        if not isinstance(date_s, str) or not DATE_SHAPE.match(date_s.strip()):
            raise IngestRejected(f"row {n}: Date {date_s!r} is not DD-MM-YYYY")
        try:
            when = datetime.strptime(date_s.strip(), DATE_FORMAT)
        except ValueError as exc:
            raise IngestRejected(f"row {n}: Date {date_s!r} is not a real DD-MM-YYYY date") from exc
        dates.add(when)
        if not isinstance(bucket, str):
            raise IngestRejected(f"row {n}: Tenor Bucket {bucket!r}")
        if isinstance(ytm, bool) or not isinstance(ytm, (int, float)):
            raise IngestRejected(f"row {n}: YTM {ytm!r} is not a number")
        value = Decimal(str(ytm))
        if not (YTM_RANGE[0] <= value <= YTM_RANGE[1]):
            raise IngestRejected(f"row {n}: YTM {value} outside {YTM_RANGE[0]}..{YTM_RANGE[1]}")
        if bucket in DROPPED:
            continue
        if bucket not in KEPT:
            raise IngestRejected(f"row {n}: unknown Tenor Bucket {bucket!r}")
        if bucket in seen:
            raise IngestRejected(f"row {n}: bucket {bucket} appears twice")
        if not isinstance(security, str) or not G_SEC.match(security.strip()):
            raise IngestRejected(f"row {n}: {bucket} holds {security!r}, not a central-government G-sec")
        seen.add(bucket)
        out.append((code_for("IN", KEPT[bucket]), when, value.quantize(QUANT), security.strip()))
    if len(dates) != 1:
        raise IngestRejected(f"expected one Date across the file, got {len(dates)}")
    missing = sorted(set(KEPT) - seen)
    if missing:
        raise IngestRejected(f"kept bucket(s) missing: {', '.join(missing)}")
    return out


def ingest(db: Session, body: bytes, uploader: str) -> dict:
    """POST /api/ingest/ccil-tenorwise. All or nothing; idempotent; one
    ingest_provenance row per accepted upload. A changed value for a date
    already held is overwritten and counted in rows_updated."""
    from database.models import IngestProvenance, TimeSeries
    from pipelines.fred_fetcher import ensure_metric, upsert_observations

    if len(body) > MAX_BODY_BYTES:
        raise IngestRejected(f"upload is {len(body)} B, over the {MAX_BODY_BYTES} B cap")
    rows = parse_ccil(body)
    try:
        ids = {m["code"]: ensure_metric(db, m).id for m in METRICS}
        inserted = updated = unchanged = 0
        for code, when, value, _security in rows:
            metric_id = ids[code]
            have = {
                (d.date() if hasattr(d, "date") else d): Decimal(str(v))
                for d, v in db.query(TimeSeries.date, TimeSeries.value)
                .filter(TimeSeries.metric_id == metric_id, TimeSeries.country_id.is_(None))
                .all()
            }
            old = have.get(when.date())
            if old is not None and old.quantize(QUANT) == value:
                unchanged += 1
                continue
            i, u = upsert_observations(db, metric_id, [(when, value)])
            inserted += i
            updated += u
        when = rows[0][1]
        db.add(IngestProvenance(
            source=INGEST_SOURCE,
            uploaded_at=datetime.utcnow(),
            uploader=(uploader or "unknown")[:100],
            sha256=hashlib.sha256(body).hexdigest(),
            rows_inserted=inserted,
            rows_updated=updated,
            rows_unchanged=unchanged,
            min_date=when,
            max_date=when,
        ))
        db.commit()
    except Exception:
        db.rollback()
        raise
    logger.info("CCIL tenorwise upload as of %s: %d inserted, %d updated, %d unchanged",
                when.date().isoformat(), inserted, updated, unchanged)
    return {
        "status": "success",
        "as_of": when.date().isoformat(),
        "inserted": inserted,
        "updated": updated,
        "unchanged": unchanged,
        "rows_changed": inserted + updated,
        "min_date": when.date().isoformat(),
        "max_date": when.date().isoformat(),
    }
