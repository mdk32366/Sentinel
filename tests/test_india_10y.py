"""Sentinel D-0118: India 10Y on the IND card. Display only, never scored.

  * monthly: FRED INDIRLTLT01STM (OECD), via the UI-only MONTHLY_10Y_EXTRA_CODES
    map; India stays OUT of SOVEREIGN_YIELD_CODES (T-0161..T-0166, spec@0.1
    renumbered by @0.3);
  * daily: Matt's CCIL Tenorwise Indicative Yields download, uploaded raw to
    POST /api/ingest/ccil-tenorwise (T-0167..T-0174, spec@0.3, §3 trimmed at
    build: provenance = the existing ingest_provenance row, raw body).

The fixture is the real file (tests/fixtures/ccil/). Variants are built from
it in memory with openpyxl. No network. The UI halves (T-0161, T-0162 card,
T-0166, T-0173) are in ui/src/components/country/india10y.dom.test.jsx.
"""
import base64
import hashlib
import io
import re
import sys
import unittest
from datetime import datetime, timedelta
from pathlib import Path
from unittest import mock

from sqlalchemy import create_engine, func
from sqlalchemy.orm import sessionmaker

from database.models import Base, Metric, TimeSeries

ROOT = Path(__file__).resolve().parents[1]
FIXTURE = ROOT / "tests" / "fixtures" / "ccil" / "TenorWiseIndicativeYields_2026-10-22.xlsx"
FIXTURE_MD5 = "9704e29c7198971ceef8bc191b5299f6"
CONSTANTS = ROOT / "ui" / "src" / "lib" / "constants.js"
MARKET = ROOT / "ui" / "src" / "hooks" / "useMarketSeries.js"
GOLDEN = ROOT / "tests" / "fixtures" / "d0113_scorer_golden.json"
COMPOSITE = ROOT / "pipelines" / "composite_stress.py"
GOLDEN_SHA256 = "05e1a1d3705383332ac4ef2c09b76fe59f6ab2e0763f785e144ad83f612a3f89"
COMPOSITE_SHA256 = "72d700db24ec41384b1f1a5aab35a7097d04c477540ab4a19bcd2876c46bad23"
ROUTE = "/api/ingest/ccil-tenorwise"
ATTRIBUTION = ("Source: The Clearing Corporation of India Ltd. (CCIL), Tenorwise Indicative Yields, "
               "https://www.ccilindia.com/tenorwise-indicative-yields")
BACKEND_SOV = {
    "JPN": "IRLTLT01JPM156N", "DEU": "IRLTLT01DEM156N", "ITA": "IRLTLT01ITM156N",
    "FRA": "IRLTLT01FRM156N", "ESP": "IRLTLT01ESM156N", "GBR": "IRLTLT01GBM156N",
    "AUS": "IRLTLT01AUM156N", "CAN": "IRLTLT01CAM156N", "NLD": "IRLTLT01NLM156N",
    "NOR": "IRLTLT01NOM156N", "SWE": "IRLTLT01SEM156N", "CHE": "IRLTLT01CHM156N",
    "BEL": "IRLTLT01BEM156N", "KOR": "IRLTLT01KRM156N",
}
EXPECTED = {"FB_IN_02Y": 6.6982, "FB_IN_05Y": 7.2097, "FB_IN_10Y": 7.4553, "FB_IN_30Y": 7.8616}


def fixture():
    return FIXTURE.read_bytes()


def session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def rows_count(db):
    return db.query(func.count(TimeSeries.id)).scalar()


def provenance_count(db):
    from database.models import IngestProvenance
    return db.query(func.count(IngestProvenance.id)).scalar()


def variant(edit):
    """The real file, opened, edited by ``edit(ws)``, saved to bytes."""
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(fixture()))
    edit(wb.active)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def rows_of(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


def rewrite(ws, rows):
    ws.delete_rows(1, ws.max_row)
    for r in rows:
        ws.append(r)


def drop_bucket(bucket):
    def edit(ws):
        rewrite(ws, [r for r in rows_of(ws) if r[1] != bucket])
    return edit


def set_cell(bucket, col, value):
    def edit(ws):
        rows = rows_of(ws)
        for r in rows[1:]:
            if r[1] == bucket:
                r[col] = value
        rewrite(ws, rows)
    return edit


def stored(db):
    out = {}
    for code, d, v in db.query(Metric.code, TimeSeries.date, TimeSeries.value).join(
            TimeSeries, TimeSeries.metric_id == Metric.id).all():
        out[code] = (d.date().isoformat() if hasattr(d, "date") else str(d)[:10], float(v))
    return out


# ------------------------------------------------------------------ T-0163 / T-0174
class T0163CompositeInvariant(unittest.TestCase):
    def test_hashes(self):
        self.assertEqual(hashlib.sha256(GOLDEN.read_bytes()).hexdigest(), GOLDEN_SHA256)
        self.assertEqual(hashlib.sha256(COMPOSITE.read_bytes()).hexdigest(), COMPOSITE_SHA256)

    def test_india_rows_change_no_composite_byte_and_ind_spread_stays_null(self):
        sys.path.insert(0, str(ROOT / "tests"))
        import composite_golden_fixture as cgf
        import test_d0113_scorer_golden as g
        from pipelines import composite_stress

        def add_india(db):
            for code, value in [("INDIRLTLT01STM", 7.06)] + [(c, v) for c, v in EXPECTED.items()]:
                m = Metric(code=code, name=code, category="x", unit="%", source="fixture")
                db.add(m)
                db.flush()
                for d in (datetime(2026, 9, 1), datetime(2026, 10, 22)):
                    db.add(TimeSeries(metric_id=m.id, country_id=None, date=d, value=value))
            db.commit()
            return db

        real_build = cgf.build
        with mock.patch.object(g, "build", lambda: add_india(real_build())):
            out = g.payload()
        self.assertEqual(out, GOLDEN.read_text(encoding="utf-8"))
        db = add_india(session())
        self.assertEqual(composite_stress.get_sovereign_spread(db, "IND", 4.0),
                         {"spread_bps": None, "widening_bps": None, "score": 0})


# ------------------------------------------------------------------ T-0164 / T-0174
class T0164CodeSetsUnchanged(unittest.TestCase):
    def test_backend_and_ui_maps_equal_the_snapshot(self):
        from pipelines import composite_stress
        self.assertEqual(composite_stress.SOVEREIGN_YIELD_CODES, BACKEND_SOV)
        js = CONSTANTS.read_text(encoding="utf-8")
        sov = js[js.index("export const SOVEREIGN_YIELD_CODES"):]
        sov = sov[: sov.index("};")]
        self.assertEqual(dict(re.findall(r"([A-Z]{3}): \"(IRLTLT01\w+)\"", sov)), BACKEND_SOV)
        self.assertNotIn("IND", sov)
        self.assertNotIn("IND", composite_stress.FB_DIM4_CODES)
        self.assertEqual(set(composite_stress.FB_DIM4_CODES), {"DEU", "GBR", "JPN", "FRA"})

    def test_extra_map_and_jobs(self):
        js = CONSTANTS.read_text(encoding="utf-8")
        self.assertRegex(js, r'export const MONTHLY_10Y_EXTRA_CODES = \{ IND: "INDIRLTLT01STM" \};')
        tracked = MARKET.read_text(encoding="utf-8")
        tracked = tracked[tracked.index("export const TRACKED_CODES"):]
        tracked = tracked[: tracked.index("];")]
        self.assertNotIn("MONTHLY_10Y_EXTRA_CODES", tracked)
        from pipelines import scheduler
        self.assertEqual(len(scheduler.FOREIGN_BOND_JOBS), 13)
        codes = {m["code"] for _, _, metrics in scheduler.FOREIGN_BOND_JOBS.values() for m in metrics}
        self.assertFalse(any(c.startswith("FB_IN_") for c in codes))
        self.assertNotIn("ccil_tenorwise", scheduler.FOREIGN_BOND_JOBS)


# ------------------------------------------------------------------ T-0162 / T-0165
class T0165Freshness(unittest.TestCase):
    def checks(self):
        from pipelines.freshness_watchdog import CHECKS
        return {c["key"]: c for c in CHECKS}

    def test_india_monthly_check(self):
        c = self.checks().get("sovereign_yield_ind")
        self.assertIsNotNone(c, "sovereign_yield_ind missing")
        self.assertEqual((c["period"], c["patterns"], c["max_age_days"], c["pipelines"]),
                         ("month", ["INDIRLTLT01STM"], 75, ["FRED"]))

    def test_sovereign_yields_still_matches_exactly_its_14(self):
        c = self.checks()["sovereign_yields"]
        self.assertEqual((c["patterns"], c["max_age_days"]), (["IRLTLT01%"], 70))
        from pipelines.fred_fetcher import FRED_METRICS as METRICS
        like = re.compile("^" + c["patterns"][0].replace("%", ".*") + "$")
        matched = sorted(m["code"] for m in METRICS if like.match(m["code"]))
        self.assertEqual(matched, sorted(BACKEND_SOV.values()))
        self.assertIn("INDIRLTLT01STM", {m["code"] for m in METRICS})

    def test_t0162_fresh_at_75_stale_at_76(self):
        from pipelines.freshness_watchdog import _classify
        c = self.checks().get("sovereign_yield_ind")
        self.assertIsNotNone(c, "sovereign_yield_ind missing")
        limit = c["max_age_days"]
        self.assertEqual(_classify(75, limit), "ok")
        self.assertNotEqual(_classify(76, limit), "ok")

    def test_fred_metric_is_display_only(self):
        from pipelines.fred_fetcher import FRED_METRICS as METRICS
        m = next((m for m in METRICS if m["code"] == "INDIRLTLT01STM"), None)
        self.assertIsNotNone(m, "INDIRLTLT01STM not in FRED_METRICS")
        self.assertEqual((m["category"], m["unit"], m["source"]), ("sovereign_yield_display", "%", "FRED"))

    def test_t0173_daily_check(self):
        c = self.checks().get("foreign_bonds_in")
        self.assertIsNotNone(c, "foreign_bonds_in missing")
        self.assertEqual((c["period"], c["max_age_days"]), ("day", 7))
        self.assertIn("A STALE badge means the upload is overdue, not an outage", c["note"])


# ------------------------------------------------------------------ T-0167
class T0167Ingest(unittest.TestCase):
    def test_fixture_identity(self):
        self.assertEqual(hashlib.md5(fixture()).hexdigest(), FIXTURE_MD5)
        self.assertEqual(len(fixture()), 3866)

    def test_exactly_four_rows_and_the_provenance_row(self):
        from database.models import IngestProvenance
        from pipelines.ccil_tenorwise import ATTRIBUTION as attr, ingest
        db = session()
        res = ingest(db, fixture(), "matt")
        self.assertEqual((res["inserted"], res["updated"], res["unchanged"], res["rows_changed"]), (4, 0, 0, 4))
        self.assertEqual(res["as_of"], "2026-10-22")
        self.assertEqual(stored(db), {c: ("2026-10-22", v) for c, v in EXPECTED.items()})
        prov = db.query(IngestProvenance).one()
        sha = hashlib.sha256(fixture()).hexdigest()
        self.assertEqual(prov.sha256, sha)
        self.assertEqual((prov.source, prov.uploader), ("ccil-tenorwise", "matt"))
        self.assertEqual((prov.rows_inserted, prov.rows_updated, prov.rows_unchanged), (4, 0, 0))
        self.assertEqual((prov.min_date.date().isoformat(), prov.max_date.date().isoformat()),
                         ("2026-10-22", "2026-10-22"))
        self.assertEqual(attr, ATTRIBUTION)  # a code constant, rendered on the card, not stored


# ------------------------------------------------------------------ T-0168
class T0168SchemaMismatch(unittest.TestCase):
    def assert_rejected(self, body):
        from pipelines.ccil_tenorwise import IngestRejected, ingest
        db = session()
        with self.assertRaises(IngestRejected):
            ingest(db, body, "matt")
        self.assertEqual((rows_count(db), provenance_count(db)), (0, 0))

    def test_renamed_header(self):
        self.assert_rejected(variant(lambda ws: ws.cell(1, 4, "YTM")))

    def test_swapped_columns(self):
        def swap(ws):
            rewrite(ws, [[r[0], r[1], r[3], r[2]] for r in rows_of(ws)])
        self.assert_rejected(variant(swap))

    def test_two_sheets_html_and_macros(self):
        self.assert_rejected(variant(lambda ws: ws.parent.create_sheet("Sheet2")))
        self.assert_rejected(b"<html><body>table</body></html>")
        import zipfile
        buf = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(fixture())) as src, zipfile.ZipFile(buf, "w") as dst:
            for item in src.infolist():
                dst.writestr(item, src.read(item))
            dst.writestr("xl/vbaProject.bin", b"\x00macro")
        self.assert_rejected(buf.getvalue())

    def test_unknown_bucket(self):
        self.assert_rejected(variant(set_cell("13Y-15Y", 1, "14Y-15Y")))


# ------------------------------------------------------------------ T-0169
class T0169Drops(unittest.TestCase):
    def test_only_the_four_g_secs_are_stored(self):
        from pipelines.ccil_tenorwise import parse_ccil
        rows = parse_ccil(fixture())
        self.assertEqual(sorted(c for c, _, _, _ in rows), sorted(EXPECTED))
        self.assertEqual({c: s for c, _, _, s in rows}["FB_IN_10Y"], "6.94% GS 2036")

    def test_missing_9y_10y_is_422_not_the_sdl_10y(self):
        T0168SchemaMismatch.assert_rejected(self, variant(drop_bucket("9Y-10Y")))

    def test_sdl_or_tbill_in_a_kept_bucket_is_422(self):
        T0168SchemaMismatch.assert_rejected(self, variant(set_cell("9Y-10Y", 2, "8.03% KARNATAKA SGS 2036")))
        T0168SchemaMismatch.assert_rejected(self, variant(set_cell("1Y-2Y", 2, "364 DTB (14/10/2027)")))

    def test_duplicate_kept_bucket_is_422(self):
        def dup(ws):
            rows = rows_of(ws)
            rewrite(ws, rows + [r for r in rows if r[1] == "4Y-5Y"])
        T0168SchemaMismatch.assert_rejected(self, variant(dup))


# ------------------------------------------------------------------ T-0170
class T0170Idempotent(unittest.TestCase):
    def test_second_upload_changes_nothing(self):
        from pipelines.ccil_tenorwise import ingest
        db = session()
        ingest(db, fixture(), "matt")
        res = ingest(db, fixture(), "matt")
        self.assertEqual((res["rows_changed"], res["unchanged"]), (0, 4))
        self.assertEqual((rows_count(db), provenance_count(db)), (4, 2))

    def test_changed_value_same_date_is_overwritten_and_logged(self):
        from database.models import IngestProvenance
        from pipelines.ccil_tenorwise import ingest
        db = session()
        ingest(db, fixture(), "matt")
        res = ingest(db, variant(set_cell("9Y-10Y", 3, 7.5)), "matt")
        self.assertEqual((res["inserted"], res["updated"], res["unchanged"], res["rows_changed"]), (0, 1, 3, 1))
        self.assertEqual(stored(db)["FB_IN_10Y"], ("2026-10-22", 7.5))
        last = db.query(IngestProvenance).order_by(IngestProvenance.id.desc()).first()
        self.assertEqual(last.rows_updated, 1)


# ------------------------------------------------------------------ T-0172
class T0172Guards(unittest.TestCase):
    def test_iso_date_mixed_dates_bad_ytm(self):
        reject = lambda b: T0168SchemaMismatch.assert_rejected(self, b)  # noqa: E731

        def all_dates(value):
            def edit(ws):
                rows = rows_of(ws)
                for r in rows[1:]:
                    r[0] = value
                rewrite(ws, rows)
            return edit
        reject(variant(all_dates("2026-10-22")))
        reject(variant(set_cell("4Y-5Y", 0, "21-10-2026")))
        reject(variant(set_cell("9Y-10Y", 3, "n/a")))
        reject(variant(set_cell("9Y-10Y", 3, 30)))
        reject(variant(set_cell("9Y-10Y", 3, -0.1)))


# ------------------------------------------------------------------ T-0171
class T0171Route(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        from sqlalchemy.pool import StaticPool
        from database.connection import get_db
        from main import app
        engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
        Base.metadata.create_all(engine)
        self.Session = sessionmaker(bind=engine)

        def _db():
            db = self.Session()
            try:
                yield db
            finally:
                db.close()

        app.dependency_overrides[get_db] = _db
        self.addCleanup(app.dependency_overrides.pop, get_db, None)
        self.app = app
        self.assertIn(ROUTE, {getattr(r, "path", None) for r in app.routes}, "CCIL ingest route not registered")

    async def send(self, body, headers=None):
        import httpx
        transport = httpx.ASGITransport(app=self.app, raise_app_exceptions=False)
        async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
            return await client.post(ROUTE, content=body, headers=headers or {})

    def basic(self, password=None):
        from config import settings
        raw = f"{settings.auth_username}:{password or settings.auth_password}".encode()
        return {"Authorization": "Basic " + base64.b64encode(raw).decode(),
                "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"}

    async def test_auth_status_codes_and_no_password_leak(self):
        from config import settings
        self.assertEqual((await self.send(fixture())).status_code, 401)
        self.assertIn((await self.send(fixture(), self.basic("wrong-password"))).status_code, (401, 403))
        with self.assertLogs(level="DEBUG") as logs:
            import logging
            logging.getLogger("sentinel.test").debug("capture")
            r = await self.send(fixture(), self.basic())
        self.assertEqual(r.status_code, 200, r.text)
        self.assertEqual(r.json()["rows_changed"], 4)
        self.assertNotIn(settings.auth_password, "\n".join(logs.output))
        self.assertNotIn(settings.auth_password, r.text)
        self.assertEqual((await self.send(b"<html>", self.basic())).status_code, 422)
        from pipelines.ccil_tenorwise import MAX_BODY_BYTES
        self.assertEqual((await self.send(b"x" * (MAX_BODY_BYTES + 1), self.basic())).status_code, 413)
        from database.models import IngestProvenance
        with self.Session() as db:
            self.assertEqual(rows_count(db), 4)
            for p in db.query(IngestProvenance).all():
                self.assertNotIn(settings.auth_password, " ".join(str(v) for v in vars(p).values()))

    async def test_proof_cookie_is_403(self):
        import httpx
        from api import proof_access
        from config import settings
        from tests.test_proof_access import _iso, _new_token, _sha, _utcnow
        proof_access._reset_lockout()
        self.addCleanup(proof_access._reset_lockout)
        token = _new_token()
        with mock.patch.object(settings, "sentinel_proof_token_sha256", _sha(token)), \
                mock.patch.object(settings, "sentinel_proof_expires_at", _iso(_utcnow() + timedelta(hours=1))):
            transport = httpx.ASGITransport(app=self.app, raise_app_exceptions=False)
            async with httpx.AsyncClient(transport=transport, base_url="http://test") as client:
                grant = await client.post("/__proof", content=f"token={token}", headers={
                    "Fly-Forwarded-Proto": "https", "Content-Type": "application/x-www-form-urlencoded"})
            cookie = grant.headers["set-cookie"].split(";", 1)[0]
            r = await self.send(fixture(), {"Cookie": cookie})
        self.assertEqual(r.status_code, 403)
        with self.Session() as db:
            self.assertEqual(rows_count(db), 0)


if __name__ == "__main__":
    unittest.main()
