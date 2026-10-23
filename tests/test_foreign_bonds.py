"""Sentinel D-0111 (vault:D-0057): Foreign Bonds - DE / UK / JP fetchers, FR option A.

Spec@0.2 §8. Every test here reads a fixture captured from the real source on
2026-10-21 (tests/fixtures/foreign_bonds/PROVENANCE.md). Nothing reaches a
network: the fetchers take an ``opener`` and a ``sleep`` so the transport,
retries and back-off run for real against canned bytes.

Each guard is written so it can go red: the negatives (header-only retry,
0-row rejection, schema drift, stale beyond the window, a Cloudflare block)
assert the failure, and the paired positives assert the same path succeeds on
good input - a guard that failed unconditionally would pass the negative alone.
"""
import io
import re
import unittest
import urllib.error
import zipfile
from datetime import datetime, timedelta
from decimal import Decimal
from email.message import Message
from pathlib import Path

from sqlalchemy import create_engine, func
from sqlalchemy.orm import sessionmaker

from database.models import Base, Metric, TimeSeries, UpdateLog

ROOT = Path(__file__).resolve().parents[1]
FIX = ROOT / "tests" / "fixtures" / "foreign_bonds"
CONSTANTS = ROOT / "ui" / "src" / "lib" / "constants.js"
MARKET_SERIES = ROOT / "ui" / "src" / "hooks" / "useMarketSeries.js"
V2_SCORER = ROOT / "pipelines" / "stress_score_v2.py"
COMPOSITE_SCORER = ROOT / "pipelines" / "composite_stress.py"

TENORS = (2, 5, 7, 10, 30)


def fixture(name):
    return (FIX / name).read_bytes()


def session():
    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    return sessionmaker(bind=engine)()


def boe_zip(xlsx_bytes, member="GLC Nominal daily data current month.xlsx"):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr(member, xlsx_bytes)
    return buf.getvalue()


class _Resp:
    def __init__(self, body):
        self.body = body

    def __enter__(self):
        return self

    def __exit__(self, *a):
        return False

    def read(self):
        return self.body


class Opener:
    """``routes``: list of (url substring, [body | exception, ...]). Each call
    pops the next item for the first matching route; the last item repeats."""

    def __init__(self, routes):
        self.routes = [(k, list(v)) for k, v in routes]
        self.calls = []

    def __call__(self, url, headers=None):
        self.calls.append(url)
        for key, items in self.routes:
            if key in url:
                item = items.pop(0) if len(items) > 1 else items[0]
                if isinstance(item, BaseException):
                    raise item
                return _Resp(item)
        raise AssertionError(f"unexpected URL {url}")


class Sleeps(list):
    def __call__(self, s):
        self.append(s)


def cloudflare_403(url):
    hdrs = Message()
    hdrs["Server"] = "cloudflare"
    body = io.BytesIO(b"<html><title>Attention Required! | Cloudflare</title>Sorry, you have been blocked</html>")
    return urllib.error.HTTPError(url, 403, "Forbidden", hdrs, body)


def ts_count(db):
    return db.query(func.count(TimeSeries.id)).scalar()


def bundesbank_ok():
    return [(f"R{t:02d}XX", [fixture(f"bundesbank_R{t:02d}XX.csv")]) for t in TENORS]


# --------------------------------------------------------------------------- parse
class ParseExactTuples(unittest.TestCase):
    def test_bundesbank_r10xx_fixture(self):
        from pipelines.bundesbank_yields import parse_bundesbank
        rows = parse_bundesbank(fixture("bundesbank_R10XX.csv"), 10)
        # 10 observations; the two weekend '.' cells are skipped, not zero.
        self.assertEqual(len(rows), 8)
        self.assertEqual(rows[0], ("FB_DE_10Y", datetime(2026, 10, 12), Decimal("3.53")))
        self.assertEqual(rows[-1], ("FB_DE_10Y", datetime(2026, 10, 21), Decimal("3.47")))
        self.assertNotIn(datetime(2026, 10, 17), [r[1] for r in rows])

    def test_bundesbank_every_tenor_latest(self):
        from pipelines.bundesbank_yields import parse_bundesbank
        latest = {t: parse_bundesbank(fixture(f"bundesbank_R{t:02d}XX.csv"), t)[-1] for t in TENORS}
        self.assertEqual(
            {t: (r[0], r[1].date().isoformat(), str(r[2].normalize())) for t, r in latest.items()},
            {2: ("FB_DE_02Y", "2026-10-21", "3.08"), 5: ("FB_DE_05Y", "2026-10-21", "3.19"),
             7: ("FB_DE_07Y", "2026-10-21", "3.31"), 10: ("FB_DE_10Y", "2026-10-21", "3.47"),
             30: ("FB_DE_30Y", "2026-10-21", "3.86")},
        )

    def test_bundesbank_refuses_a_series_for_the_wrong_tenor(self):
        from pipelines.bundesbank_yields import parse_bundesbank
        from pipelines.foreign_bond_common import SchemaDrift
        with self.assertRaises(SchemaDrift):
            parse_bundesbank(fixture("bundesbank_R10XX.csv"), 5)

    def test_boe_current_month_fixture(self):
        from pipelines.boe_yields import parse_boe_xlsx
        rows = parse_boe_xlsx(fixture("boe_glc_nominal_current_month.xlsx"))
        by = {(c, d.date().isoformat()): v for c, d, v in rows}
        self.assertEqual(by[("FB_GB_10Y", "2026-10-20")], Decimal("5.38407521"))
        self.assertEqual(by[("FB_GB_02Y", "2026-10-20")], Decimal("4.72034099"))
        self.assertEqual(by[("FB_GB_30Y", "2026-10-20")], Decimal("5.85337290"))
        # 14 business days 2026-10-01..20, five tenors each.
        self.assertEqual(len(rows), 70)
        self.assertEqual(max(d for _, d, _ in rows), datetime(2026, 10, 20))

    def test_boe_archive_fixture_skips_banner_and_holiday_rows(self):
        from pipelines.boe_yields import parse_boe_xlsx
        rows = parse_boe_xlsx(fixture("boe_glc_nominal_archive_2025_trimmed.xlsx"))
        dates = sorted({d.date().isoformat() for _, d, _ in rows})
        # 2025-01-01 is in the file with no values; the 'Refresh' row is a banner.
        self.assertEqual(dates, ["2025-01-02", "2025-01-03", "2026-09-28", "2026-09-29", "2026-09-30"])
        self.assertEqual(len(rows), 25)

    def test_boe_columns_are_found_by_years_value_not_index(self):
        import openpyxl
        from pipelines.boe_yields import parse_boe_xlsx
        wb = openpyxl.load_workbook(io.BytesIO(fixture("boe_glc_nominal_current_month.xlsx")))
        ws = wb["4. spot curve"]
        ws.insert_cols(2)  # shift every maturity one column right
        buf = io.BytesIO()
        wb.save(buf)
        moved = {(c, d): v for c, d, v in parse_boe_xlsx(buf.getvalue())}
        orig = {(c, d): v for c, d, v in parse_boe_xlsx(fixture("boe_glc_nominal_current_month.xlsx"))}
        self.assertEqual(moved, orig)

    def test_mof_current_month_fixture(self):
        from pipelines.mof_jgb import parse_mof_csv
        rows = parse_mof_csv(fixture("mof_jgbcm.csv"))
        by = {(c, d.date().isoformat()): v for c, d, v in rows}
        self.assertEqual(by[("FB_JP_10Y", "2026-10-20")], Decimal("2.927"))
        self.assertEqual(by[("FB_JP_02Y", "2026-10-20")], Decimal("1.909"))
        self.assertEqual(by[("FB_JP_05Y", "2026-10-20")], Decimal("2.313"))
        self.assertEqual(by[("FB_JP_07Y", "2026-10-20")], Decimal("2.516"))
        self.assertEqual(by[("FB_JP_30Y", "2026-10-20")], Decimal("3.992"))
        self.assertEqual(len(rows), 65)  # 13 dates x 5

    def test_mof_era_dates(self):
        from pipelines.foreign_bond_common import SchemaDrift
        from pipelines.mof_jgb import era_to_date
        self.assertEqual(era_to_date("R8.10.20"), datetime(2026, 10, 20))
        self.assertEqual(era_to_date("R1.5.7"), datetime(2019, 5, 7))
        self.assertEqual(era_to_date("H30.4.2"), datetime(2018, 4, 2))
        self.assertEqual(era_to_date("S49.9.24"), datetime(1974, 9, 24))
        for bad in ("2026-10-20", "X8.10.20", "R8.13.1", ""):
            with self.subTest(bad=bad), self.assertRaises(SchemaDrift):
                era_to_date(bad)

    def test_mof_history_fixture_skips_dash_cells(self):
        from pipelines.mof_jgb import parse_mof_csv
        rows = parse_mof_csv(fixture("mof_jgbcm_all_trimmed.csv"))
        first = [r for r in rows if r[1] == datetime(1974, 9, 24)]
        # 1974: 2Y/5Y/7Y published, 10Y and 30Y are '-' -> absent, not zero.
        self.assertEqual(sorted(r[0] for r in first), ["FB_JP_02Y", "FB_JP_05Y", "FB_JP_07Y"])
        self.assertEqual(rows[-1][1], datetime(2026, 9, 30))


# --------------------------------------------------------------------- schema drift
class SchemaDriftFailsClosed(unittest.TestCase):
    def _assert_failed_and_nothing_written(self, run, **kw):
        db = session()
        before = ts_count(db)
        result = run(db, **kw)
        self.assertEqual(result["status"], "failed")
        self.assertEqual(ts_count(db), before)
        self.assertEqual(ts_count(db), 0)
        log = db.query(UpdateLog).one()
        self.assertEqual(log.status, "failed")
        self.assertIn("FETCH_PARSE", log.error_message)
        return result

    def test_bundesbank_column_renamed(self):
        from pipelines.bundesbank_yields import run_bundesbank_yields
        bad = fixture("bundesbank_R10XX.csv").replace(b";OBS_VALUE;", b";OBS_VAL;", 1)
        routes = bundesbank_ok()
        routes[3] = ("R10XX", [bad])
        self._assert_failed_and_nothing_written(
            run_bundesbank_yields, opener=Opener(routes), sleep=Sleeps())

    def test_bundesbank_value_99_rejects_the_whole_source(self):
        # One bad cell in the 30Y file: the 02..10Y files parsed fine and must
        # still not be written (all-or-nothing per run).
        from pipelines.bundesbank_yields import run_bundesbank_yields
        bad = fixture("bundesbank_R30XX.csv").replace(b";2026-10-21;3.86;", b";2026-10-21;99;", 1)
        self.assertNotEqual(bad, fixture("bundesbank_R30XX.csv"))
        routes = bundesbank_ok()
        routes[4] = ("R30XX", [bad])
        self._assert_failed_and_nothing_written(
            run_bundesbank_yields, opener=Opener(routes), sleep=Sleeps())

    def test_boe_years_row_missing_30(self):
        import openpyxl
        from pipelines.boe_yields import run_boe_yields
        wb = openpyxl.load_workbook(io.BytesIO(fixture("boe_glc_nominal_current_month.xlsx")))
        ws = wb["4. spot curve"]
        for row in ws.iter_rows():
            if row[0].value == "years:":
                for cell in row:
                    if cell.value == 30:
                        cell.value = 29.5
        buf = io.BytesIO()
        wb.save(buf)
        self._assert_failed_and_nothing_written(
            run_boe_yields, opener=Opener([("latest-yield-curve-data.zip", [boe_zip(buf.getvalue())])]),
            sleep=Sleeps())

    def test_mof_date_header_changed(self):
        from pipelines.mof_jgb import run_mof_jgb
        bad = fixture("mof_jgbcm.csv").replace("基準日".encode("cp932"), "日付".encode("cp932"), 1)
        self._assert_failed_and_nothing_written(
            run_mof_jgb, opener=Opener([("jgbcm.csv", [bad])]), sleep=Sleeps())

    def test_mof_value_99(self):
        from pipelines.mof_jgb import run_mof_jgb
        bad = fixture("mof_jgbcm.csv").replace(b",2.927,", b",99,", 1)
        self.assertNotEqual(bad, fixture("mof_jgbcm.csv"))
        self._assert_failed_and_nothing_written(
            run_mof_jgb, opener=Opener([("jgbcm.csv", [bad])]), sleep=Sleeps())

    def test_mof_unparseable_date(self):
        from pipelines.mof_jgb import run_mof_jgb
        bad = fixture("mof_jgbcm.csv").replace(b"R8.10.20,", b"2026/10/20,", 1)
        self._assert_failed_and_nothing_written(
            run_mof_jgb, opener=Opener([("jgbcm.csv", [bad])]), sleep=Sleeps())


# ------------------------------------------------------------- header-only / 0 rows
class HeaderOnlyIsRetriedThenFailed(unittest.TestCase):
    def test_header_only_every_time_is_failed_never_a_0_row_success(self):
        from pipelines.bundesbank_yields import run_bundesbank_yields
        routes = bundesbank_ok()
        routes[2] = ("R07XX", [fixture("bundesbank_header_only.csv")])
        opener, sleeps = Opener(routes), Sleeps()
        db = session()
        result = run_bundesbank_yields(db, opener=opener, sleep=sleeps)
        self.assertEqual(result["status"], "failed")
        self.assertNotEqual(result["status"], "success")
        self.assertEqual(ts_count(db), 0)
        # Retried: three attempts at R07XX, backoff 4 s then 8 s.
        self.assertEqual(sum("R07XX" in u for u in opener.calls), 3)
        self.assertEqual(list(sleeps), [4.0, 8.0])
        log = db.query(UpdateLog).one()
        self.assertEqual((log.pipeline_name, log.status, log.records_inserted),
                         ("Bundesbank_Yields", "failed", 0))
        self.assertIn("header", log.error_message)

    def test_header_only_once_then_data_succeeds(self):
        # The observed 10/21 quirk: R07XX header-only once, fine on retry.
        from pipelines.bundesbank_yields import run_bundesbank_yields
        routes = bundesbank_ok()
        routes[2] = ("R07XX", [fixture("bundesbank_header_only.csv"), fixture("bundesbank_R07XX.csv")])
        opener, sleeps = Opener(routes), Sleeps()
        db = session()
        result = run_bundesbank_yields(db, opener=opener, sleep=sleeps)
        self.assertEqual(result["status"], "success")
        self.assertEqual(list(sleeps), [4.0])
        self.assertEqual(result["inserted"], 40)  # 5 tenors x 8 observations
        self.assertEqual(ts_count(db), 40)
        self.assertEqual(result["latest_date"], "2026-10-21")

    def test_is_header_only_detector(self):
        from pipelines.bundesbank_yields import is_header_only
        self.assertTrue(is_header_only(fixture("bundesbank_header_only.csv")))
        self.assertTrue(is_header_only(fixture("bundesbank_header_only.csv") + b"\n\n"))
        self.assertFalse(is_header_only(fixture("bundesbank_R10XX.csv")))


class ZeroRowsIsNeverSuccess(unittest.TestCase):
    def test_mof_header_without_rows_is_failed(self):
        from pipelines.mof_jgb import run_mof_jgb
        lines = fixture("mof_jgbcm.csv").split(b"\r\n") if b"\r\n" in fixture("mof_jgbcm.csv") \
            else fixture("mof_jgbcm.csv").split(b"\n")
        header_only = b"\n".join(lines[:2]) + b"\n"
        db = session()
        result = run_mof_jgb(db, opener=Opener([("jgbcm.csv", [header_only])]), sleep=Sleeps())
        self.assertEqual(result["status"], "failed")
        self.assertEqual(ts_count(db), 0)
        self.assertEqual(db.query(UpdateLog).one().status, "failed")

    def test_boe_month_start_with_no_rows_is_failed(self):
        import openpyxl
        from pipelines.boe_yields import run_boe_yields
        wb = openpyxl.load_workbook(io.BytesIO(fixture("boe_glc_nominal_current_month.xlsx")))
        ws = wb["4. spot curve"]
        for row in ws.iter_rows():
            if isinstance(row[0].value, datetime):
                for cell in row:
                    cell.value = None
        buf = io.BytesIO()
        wb.save(buf)
        db = session()
        result = run_boe_yields(db, opener=Opener([("latest-yield-curve-data.zip", [boe_zip(buf.getvalue())])]),
                                sleep=Sleeps())
        self.assertEqual(result["status"], "failed")
        self.assertEqual(ts_count(db), 0)


# ------------------------------------------------------------------------- blocked
class BlockedIsClassified(unittest.TestCase):
    def test_cloudflare_403_is_blocked_and_not_retried(self):
        from pipelines.fetch_failure import ClassifiedFetchError
        from pipelines.foreign_bond_common import fetch_bytes
        url = "https://www.mof.go.jp/jgbs/reference/interest_rate/jgbcm.csv"
        sleeps = Sleeps()
        with self.assertRaises(ClassifiedFetchError) as ctx:
            fetch_bytes(url, opener=Opener([("jgbcm", [cloudflare_403(url)])]), sleep=sleeps)
        self.assertEqual(ctx.exception.failure.kind, "blocked")
        self.assertTrue(ctx.exception.failure.cloudflare)
        self.assertEqual(list(sleeps), [])

    def test_a_blocked_run_logs_fetch_blocked(self):
        from pipelines.mof_jgb import run_mof_jgb
        url = "https://www.mof.go.jp/jgbs/reference/interest_rate/jgbcm.csv"
        db = session()
        result = run_mof_jgb(db, opener=Opener([("jgbcm", [cloudflare_403(url)])]), sleep=Sleeps())
        self.assertEqual(result["status"], "failed")
        msg = db.query(UpdateLog).one().error_message
        self.assertTrue(msg.startswith("FETCH_BLOCKED"), msg)
        self.assertIn("cf=1", msg)
        self.assertIn("host=www.mof.go.jp", msg)

    def test_a_5xx_is_retried_three_times(self):
        from pipelines.foreign_bond_common import fetch_bytes
        url = "https://www.bankofengland.co.uk/x.zip"
        err = urllib.error.HTTPError(url, 503, "Unavailable", Message(), io.BytesIO(b""))
        opener, sleeps = Opener([("x.zip", [err])]), Sleeps()
        with self.assertRaises(Exception):
            fetch_bytes(url, opener=opener, sleep=sleeps)
        self.assertEqual(len(opener.calls), 3)
        self.assertEqual(list(sleeps), [4.0, 8.0])


# ---------------------------------------------------------------------- idempotent
class IdempotentRuns(unittest.TestCase):
    def test_every_source_twice_gives_no_duplicates(self):
        from pipelines.boe_yields import run_boe_yields
        from pipelines.bundesbank_yields import run_bundesbank_yields
        from pipelines.mof_jgb import run_mof_jgb
        db = session()
        runs = [
            (run_bundesbank_yields, lambda: Opener(bundesbank_ok()), 40),
            (run_boe_yields, lambda: Opener([("latest-yield-curve-data.zip",
                                              [boe_zip(fixture("boe_glc_nominal_current_month.xlsx"))])]), 70),
            (run_mof_jgb, lambda: Opener([("jgbcm.csv", [fixture("mof_jgbcm.csv")])]), 65),
        ]
        total = 0
        for run, opener, n in runs:
            with self.subTest(run=run.__name__):
                first = run(db, opener=opener(), sleep=Sleeps())
                second = run(db, opener=opener(), sleep=Sleeps())
                total += n
                self.assertEqual((first["status"], first["inserted"], first["updated"]), ("success", n, 0))
                self.assertEqual((second["status"], second["inserted"], second["updated"]), ("success", 0, n))
                self.assertEqual(ts_count(db), total)
        self.assertEqual(db.query(func.count(Metric.id)).scalar(), 15)
        cats = {m.category for m in db.query(Metric).all()}
        self.assertEqual(cats, {"foreign_bond"})


# ------------------------------------------------------------------------- stale
class FreshnessWindows(unittest.TestCase):
    """DE 4 d, UK 5 d, JP 7 d (spec §5). One day inside passes, one outside trips."""

    CASES = (("foreign_bonds_de", "FB_DE_10Y", 4), ("foreign_bonds_gb", "FB_GB_10Y", 5),
             ("foreign_bonds_jp", "FB_JP_10Y", 7))

    def _status(self, code, key, days_ago):
        from pipelines.freshness_watchdog import get_freshness_report
        db = session()
        m = Metric(code=code, name=code, category="foreign_bond", unit="%", source="t")
        db.add(m)
        db.commit()
        db.add(TimeSeries(metric_id=m.id, country_id=None,
                          date=datetime.utcnow() - timedelta(days=days_ago), value=3.0))
        db.commit()
        src = next(s for s in get_freshness_report(db)["sources"] if s["key"] == key)
        return src

    def test_windows(self):
        for key, code, window in self.CASES:
            with self.subTest(key=key):
                inside = self._status(code, key, window - 1)
                outside = self._status(code, key, window + 1)
                self.assertEqual(inside["max_age_days"], window)
                self.assertEqual(inside["status"], "ok")
                self.assertEqual(outside["status"], "stale")

    def test_the_spec_ages(self):
        # Spec §8 verbatim: DE 5 vs 3, GB 6 vs 4, JP 8 vs 6.
        for (key, code, _), (bad, good) in zip(self.CASES, ((5, 3), (6, 4), (8, 6))):
            with self.subTest(key=key):
                self.assertEqual(self._status(code, key, bad)["status"], "stale")
                self.assertEqual(self._status(code, key, good)["status"], "ok")

    def test_france_stays_on_the_monthly_oecd_source(self):
        from pipelines.freshness_watchdog import CHECKS
        sov = next(c for c in CHECKS if c["key"] == "sovereign_yields")
        self.assertEqual((sov["patterns"], sov["max_age_days"], sov["period"]),
                         (["IRLTLT01%"], 70, "month"))


# ----------------------------------------------------------------------- isolation
def _like(pattern, code):
    rx = ""
    i = 0
    while i < len(pattern):
        ch = pattern[i]
        if ch == "\\" and i + 1 < len(pattern):
            rx += re.escape(pattern[i + 1])
            i += 2
            continue
        rx += ".*" if ch == "%" else "." if ch == "_" else re.escape(ch)
        i += 1
    return re.fullmatch(rx, code) is not None


FB_CODES = [f"FB_{c}_{t:02d}Y" for c in ("DE", "GB", "JP") for t in TENORS]
# Sentinel D-0112: France joins, daily, from the Banque de France. Every
# isolation guard below runs over all twenty codes.
FB_CODES_ALL = FB_CODES + [f"FB_FR_{t:02d}Y" for t in TENORS]
# Sentinel D-0115 (batch 2a): each source at its own tenors, plus the ECB
# euro-area benchmark curves. The same isolation guards run over all of them.
BATCH2A_TENORS = {"CA": (2, 3, 5, 7, 10, 30), "SE": (2, 5, 7, 10), "AU": (2, 3, 5, 10),
                  "NO": (3, 5, 7, 10), "ES": (3, 5, 10, 15),
                  "EA_AAA": TENORS, "EA_ALL": TENORS}
FB_CODES_2A = [f"FB_{c}_{t:02d}Y" for c, ts in BATCH2A_TENORS.items() for t in ts]
FB_CODES_ALL = FB_CODES_ALL + FB_CODES_2A
# Sentinel D-0116 (batch 2b-1): PL auctions, IT BMK, and the ECB IRS monthly 10Y.
FB_CODES_2B1 = ([f"FB_PL_{t:02d}Y" for t in (2, 5, 10)] + [f"FB_IT_{t:02d}Y" for t in (3, 5, 10, 30)]
                + [f"FB_{c}_10Y_ECB_M" for c in ("IT", "ES", "GR", "PT", "PL")])
FB_CODES_ALL = FB_CODES_ALL + FB_CODES_2B1


class IsolationFromScoring(unittest.TestCase):
    def test_fifteen_codes_exactly(self):
        from pipelines.boe_yields import METRICS as GB
        from pipelines.bundesbank_yields import METRICS as DE
        from pipelines.mof_jgb import METRICS as JP
        codes = [m["code"] for m in DE + GB + JP]
        self.assertEqual(codes, FB_CODES)
        for m in DE + GB + JP:
            self.assertLessEqual(len(m["code"]), 50)
            self.assertLessEqual(len(m["description"]), 500)
            self.assertEqual(m["category"], "foreign_bond")
            self.assertIn("D-0111", m["description"])

    def test_no_fb_code_matches_a_scored_pattern(self):
        for pattern in ("IRLTLT01%", "%\\_CDS\\_%"):
            for code in FB_CODES_ALL:
                self.assertFalse(_like(pattern, code), f"{code} LIKE {pattern}")
        # the matcher itself can go red
        self.assertTrue(_like("IRLTLT01%", "IRLTLT01FRM156N"))
        self.assertTrue(_like("%\\_CDS\\_%", "RUS_CDS_5Y"))

    def test_no_fb_code_in_a_scorer_or_the_default_board(self):
        js = CONSTANTS.read_text(encoding="utf-8")
        metrics_block = js[js.index("export const METRICS = ["):]
        metrics_block = metrics_block[: metrics_block.index("\n];")]
        sov = js[js.index("export const SOVEREIGN_YIELD_CODES"):]
        sov = sov[: sov.index("};")]
        tracked = MARKET_SERIES.read_text(encoding="utf-8")
        tracked = tracked[tracked.index("export const TRACKED_CODES"):]
        tracked = tracked[: tracked.index("];")]
        for code in FB_CODES_ALL + ["FB_"]:
            with self.subTest(code=code):
                self.assertNotIn(code, V2_SCORER.read_text(encoding="utf-8"))
                # D-0114 amends D-0111: the composite scorer reads exactly the
                # four FB_xx_10Y codes for dimension 4, and nothing else FB_.
                # tests/test_markets_tooltips.py pins that set.
                self.assertNotIn(code, metrics_block)
                self.assertNotIn(code, sov)
                self.assertNotIn(code, tracked)
        self.assertNotIn("foreign_bond", V2_SCORER.read_text(encoding="utf-8"))
        self.assertNotIn('"foreign_bond"', COMPOSITE_SCORER.read_text(encoding="utf-8"))

    def test_the_watchdog_patterns_do_not_overlap(self):
        from pipelines.freshness_watchdog import CHECKS
        fb = [c for c in CHECKS if c["key"].startswith("foreign_bonds_")]
        self.assertEqual(sorted(c["key"] for c in fb),
                         ["foreign_bonds_au", "foreign_bonds_ca", "foreign_bonds_de",
                          "foreign_bonds_ea", "foreign_bonds_es", "foreign_bonds_fr",
                          "foreign_bonds_gb", "foreign_bonds_it", "foreign_bonds_jp",
                          "foreign_bonds_no", "foreign_bonds_pl", "foreign_bonds_se"])
        # D-0116: the ECB IRS monthly fallbacks have their own check, outside
        # the per-country prefix, and own exactly the five _ECB_M codes.
        irs = [c for c in CHECKS if c["key"] == "ecb_irs_monthly"]
        self.assertEqual(len(irs), 1)
        for c in CHECKS:
            for p in c["patterns"]:
                owned = [code for code in FB_CODES_ALL if _like(p, code)]
                if c["key"].startswith("foreign_bonds_"):
                    cc = c["key"][-2:].upper()
                    mine = [code for code in FB_CODES_ALL
                            if code.startswith(f"FB_{cc}_") and not code.endswith("_ECB_M")]
                    self.assertEqual(owned, mine, (c["key"], p))
                    # PL has three tenors (D-0116); every other country four or more.
                    self.assertGreaterEqual(len(owned), 3, (c["key"], p))
                elif c["key"] == "ecb_irs_monthly":
                    self.assertEqual(owned, [code for code in FB_CODES_ALL if code.endswith("_ECB_M")])
                else:
                    self.assertEqual(owned, [], (c["key"], p))


# ----------------------------------------------------------------------- scheduler
class Scheduled(unittest.TestCase):
    NAMES = {"Bundesbank_Yields", "BoE_Yields", "MOF_JGB", "BdF_TEC",
             # Sentinel D-0115 (batch 2a).
             "BoC_Valet", "Riksbank_SWEA", "RBA_F2", "Norges_GOVT", "BdE_TI13", "ECB_YC",
             # Sentinel D-0116 (batch 2b-1).
             "MinFin_PL", "BdI_BMK", "ECB_IRS"}

    def test_names_declared(self):
        from pipelines import scheduler as sched
        self.assertLessEqual(self.NAMES, sched.SCHEDULED_PIPELINES)

    def test_watchdog_pipelines(self):
        from pipelines.freshness_watchdog import CHECKS
        got = {c["key"]: (c["pipelines"], c["max_age_days"], c["period"], c["patterns"])
               for c in CHECKS if c["key"].startswith("foreign_bonds_")}
        self.assertEqual(got, {
            "foreign_bonds_de": (["Bundesbank_Yields"], 4, "day", ["FB\\_DE\\_%"]),
            "foreign_bonds_gb": (["BoE_Yields"], 5, "day", ["FB\\_GB\\_%"]),
            "foreign_bonds_jp": (["MOF_JGB"], 7, "day", ["FB\\_JP\\_%"]),
            "foreign_bonds_fr": (["BdF_TEC"], 4, "day", ["FB\\_FR\\_%"]),
            # Sentinel D-0115 spec@0.3 §15 limits (T-0118).
            "foreign_bonds_ca": (["BoC_Valet"], 4, "day", ["FB\\_CA\\_%"]),
            "foreign_bonds_se": (["Riksbank_SWEA"], 4, "day", ["FB\\_SE\\_%"]),
            "foreign_bonds_au": (["RBA_F2"], 12, "day", ["FB\\_AU\\_%"]),
            "foreign_bonds_no": (["Norges_GOVT"], 4, "day", ["FB\\_NO\\_%"]),
            # D-0116: %Y so the ES card check does not own FB_ES_10Y_ECB_M.
            "foreign_bonds_es": (["BdE_TI13"], 4, "day", ["FB\\_ES\\_%Y"]),
            "foreign_bonds_ea": (["ECB_YC"], 4, "day", ["FB\\_EA\\_%"]),
            # Sentinel D-0116 (batch 2b-1).
            "foreign_bonds_pl": (["MinFin_PL"], 40, "day", ["FB\\_PL\\_%Y"]),
            "foreign_bonds_it": (["BdI_BMK"], 55, "month", ["FB\\_IT\\_%Y"]),
        })

    def test_declared_by_the_modules(self):
        from pipelines.freshness_watchdog import declared_pipeline_names
        self.assertLessEqual(self.NAMES, declared_pipeline_names())


if __name__ == "__main__":
    unittest.main()
